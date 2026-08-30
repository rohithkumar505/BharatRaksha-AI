#!/usr/bin/env python3
"""Generate Decode SIH 2026 Ideation PPT for BharatRaksha AI."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from copy import deepcopy
from lxml import etree

# Colors
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
SAFFRON = RGBColor(0xE6, 0x7E, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x4A, 0x55, 0x68)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x1E, 0x84, 0x4A)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
CARD = RGBColor(0xEE, 0xF3, 0xF8)


def set_run(run, size=14, bold=False, color=DARK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=NAVY):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_header_bar(slide, title):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.0)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.0), Inches(13.333), Inches(0.08)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = SAFFRON
    accent.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.22), Inches(12.5), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, 28, True, WHITE)
    brand = slide.shapes.add_textbox(Inches(10.2), Inches(0.28), Inches(2.8), Inches(0.4))
    bp = brand.text_frame.paragraphs[0]
    bp.alignment = PP_ALIGN.RIGHT
    br = bp.add_run()
    br.text = "OSCode × Decode SIH 2026"
    set_run(br, 11, False, RGBColor(0xB8, 0xC4, 0xD6))


def add_card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE8)
    shape.line.width = Pt(1)
    return shape


def add_text(slide, left, top, width, height, text, size=13, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def add_bullets(slide, left, top, width, height, items, size=12, color=DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = f"• {item}"
        set_run(run, size, False, color)
    return box


def flow_box(slide, left, top, width, height, text, fill=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, 11, True, WHITE)
    try:
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        shape.text_frame.auto_size = None
    except Exception:
        pass
    return shape


def arrow_right(slide, left, top):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, left, top, Inches(0.35), Inches(0.22)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = SAFFRON
    shape.line.fill.background()
    return shape


def arrow_down(slide, left, top):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, left, top, Inches(0.22), Inches(0.28)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = SAFFRON
    shape.line.fill.background()
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ========== SLIDE 1: COVER ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY)
    # saffron line
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(2.55), Inches(4.3), Inches(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = SAFFRON
    line.line.fill.background()

    add_text(s, Inches(1), Inches(1.5), Inches(11.3), Inches(0.4),
             "OSCode by communityX  |  DECODE SIH 2026", 14, False, RGBColor(0xB8, 0xC4, 0xD6), PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(2.7), Inches(11.3), Inches(0.8),
             "BharatRaksha AI", 48, True, WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(3.55), Inches(11.3), Inches(0.5),
             "The AI Shield for Every Disaster", 24, False, SAFFRON, PP_ALIGN.CENTER)
    add_text(s, Inches(1.5), Inches(4.4), Inches(10.3), Inches(0.8),
             "AI-Powered Disaster Response Intelligence Platform\nPredict • Respond • Rescue • Recover", 16, False, WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(5.7), Inches(11.3), Inches(0.4),
             "PRESENTED BY", 12, True, RGBColor(0xB8, 0xC4, 0xD6), PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(6.1), Inches(11.3), Inches(0.4),
             "[ Team Name ]", 20, True, WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(6.7), Inches(11.3), Inches(0.3),
             "PS3: Disaster Response Intelligence Platform  |  Software-Only Solution", 12, False, RGBColor(0x9A, 0xA8, 0xBC), PP_ALIGN.CENTER)

    # ========== SLIDE 2: TEAM ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "TEAM MEMBERS")
    roles = [
        ("Team Member 1", "Team Lead / Full Stack"),
        ("Team Member 2", "AI / ML Engineer"),
        ("Team Member 3", "Frontend / UI-UX"),
        ("Team Member 4", "Backend / Database"),
        ("Team Member 5", "Research / Pitch / Docs"),
    ]
    for i, (name, role) in enumerate(roles):
        left = Inches(0.5 + (i % 5) * 2.5)
        top = Inches(2.4)
        card = add_card(s, left, top, Inches(2.3), Inches(2.6), WHITE)
        circle = s.shapes.add_shape(
            MSO_SHAPE.OVAL, left + Inches(0.55), top + Inches(0.35), Inches(1.2), Inches(1.2)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = NAVY
        circle.line.fill.background()
        add_text(s, left, top + Inches(0.7), Inches(2.3), Inches(0.5),
                 str(i + 1), 22, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, left + Inches(0.1), top + Inches(1.7), Inches(2.1), Inches(0.4),
                 name, 13, True, NAVY, PP_ALIGN.CENTER)
        add_text(s, left + Inches(0.1), top + Inches(2.1), Inches(2.1), Inches(0.4),
                 role, 11, False, GRAY, PP_ALIGN.CENTER)
    add_text(s, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.2),
             "Note: Replace placeholder names with actual team member names, college, and contact details before final submission.",
             12, False, GRAY)

    # ========== SLIDE 3: PROBLEM STATEMENT ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "PROBLEM STATEMENT")

    add_card(s, Inches(0.4), Inches(1.3), Inches(12.5), Inches(1.3), WHITE)
    add_text(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.35),
             "Problem Title", 12, True, SAFFRON)
    add_text(s, Inches(0.6), Inches(1.75), Inches(12), Inches(0.7),
             "PS3: Disaster Response Intelligence Platform for flood prediction, emergency planning, and resource allocation.",
             16, True, NAVY)

    add_card(s, Inches(0.4), Inches(2.85), Inches(6.1), Inches(4.1), WHITE)
    add_text(s, Inches(0.6), Inches(3.0), Inches(5.7), Inches(0.35),
             "What does this problem aim to solve?", 13, True, SAFFRON)
    add_bullets(s, Inches(0.6), Inches(3.45), Inches(5.7), Inches(3.3), [
        "Delayed disaster warnings and unclear risk information",
        "People cannot find safe routes, shelters, or hospitals quickly",
        "Rescue teams lack live coordination and resource visibility",
        "Health and farming recovery support is fragmented after disasters",
        "Language and low-network barriers reduce access to help",
    ], 13)

    add_card(s, Inches(6.8), Inches(2.85), Inches(6.1), Inches(4.1), WHITE)
    add_text(s, Inches(7.0), Inches(3.0), Inches(5.7), Inches(0.35),
             "Why did our team choose this?", 13, True, SAFFRON)
    add_bullets(s, Inches(7.0), Inches(3.45), Inches(5.7), Inches(3.3), [
        "India faces recurring multi-hazard disasters every year",
        "High social impact across citizens, farmers, hospitals, and government",
        "Strong fit for Generative AI, NLP, Computer Vision, and Maps",
        "Can unify prediction, response, and recovery in one platform",
        "Aligns with Digital India and disaster-management modernization",
    ], 13)

    # ========== SLIDE 4: REAL-WORLD ALIGNMENT ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "REAL-WORLD PROBLEM ALIGNMENT")

    cards = [
        ("Who faces this issue?", [
            "Citizens in flood/cyclone-prone regions",
            "Farmers facing crop and livestock loss",
            "Hospitals and emergency medical teams",
            "NDRF / SDRF / Police / Fire & Rescue",
            "Local government and NGO volunteers",
        ]),
        ("Current gap in the system", [
            "Alerts, maps, SOS, health, and agri tools are siloed",
            "English-heavy interfaces exclude many users",
            "Weak offline support in low-connectivity areas",
            "Slow resource allocation and unclear command view",
            "Damage assessment is mostly manual and delayed",
        ]),
        ("How we fill the gap", [
            "One All-in-One multilingual AI Assistant",
            "Predict + Respond + Recover in a single product",
            "Live maps, healthcare, agri recovery, and dashboards",
            "Offline emergency mode through PWA",
            "Software-only — works on phone/web, no custom hardware",
        ]),
    ]
    for i, (title, items) in enumerate(cards):
        left = Inches(0.35 + i * 4.3)
        add_card(s, left, Inches(1.3), Inches(4.1), Inches(4.5), WHITE)
        add_text(s, left + Inches(0.2), Inches(1.45), Inches(3.7), Inches(0.5), title, 14, True, NAVY)
        add_bullets(s, left + Inches(0.2), Inches(2.05), Inches(3.7), Inches(3.5), items, 12)

    add_card(s, Inches(0.35), Inches(6.0), Inches(12.6), Inches(1.15), NAVY)
    add_text(s, Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.35),
             "Connected Schemes / Missions", 12, True, SAFFRON)
    add_text(s, Inches(0.55), Inches(6.5), Inches(12.2), Inches(0.5),
             "Digital India  •  NDMA / State Disaster Management frameworks  •  Ayushman Bharat (emergency health access guidance)  •  PM Fasal Bima / state relief & compensation guidance  •  Common Emergency Number 112 ecosystem",
             12, False, WHITE)

    # ========== SLIDE 5: PROPOSED SOLUTION ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "PROPOSED SOLUTION")

    add_card(s, Inches(0.35), Inches(1.25), Inches(12.6), Inches(2.0), WHITE)
    add_text(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(0.3),
             "What are we building? (100–150 words)", 12, True, SAFFRON)
    add_text(s, Inches(0.55), Inches(1.7), Inches(12.2), Inches(1.4),
             "BharatRaksha AI is an AI-powered Disaster Response Intelligence Platform for India. It supports citizens, rescue teams, hospitals, farmers, and government authorities before, during, and after disasters. Through one multilingual AI Assistant (chat + voice), users can check disaster risk, view live maps and safe routes, find shelters and hospitals, get first-aid and scheme guidance, report incidents, and assess crop or infrastructure damage from images. A government/rescue dashboard enables live monitoring, resource allocation, shelter occupancy tracking, and volunteer coordination. The system is fully software-based with offline emergency support — no custom hardware required.",
             13, False, DARK)

    # Key features
    feats = [
        ("All-in-One AI Assistant", "Multilingual chat + voice that can run the full platform"),
        ("Predict & Navigate", "Multi-hazard risk score, live maps, safe routes, shelters"),
        ("Health + Agri Recovery", "Emergency medical support and crop damage / compensation help"),
        ("Govt Command Dashboard", "Rescue tracking, resources, shelters, volunteers, analytics"),
    ]
    for i, (t, d) in enumerate(feats):
        left = Inches(0.35 + i * 3.2)
        add_card(s, left, Inches(3.5), Inches(3.05), Inches(2.0), WHITE)
        add_text(s, left + Inches(0.15), Inches(3.65), Inches(2.75), Inches(0.5), t, 13, True, NAVY)
        add_text(s, left + Inches(0.15), Inches(4.25), Inches(2.75), Inches(1.0), d, 11, False, GRAY)

    add_card(s, Inches(0.35), Inches(5.7), Inches(12.6), Inches(1.45), NAVY)
    add_text(s, Inches(0.55), Inches(5.85), Inches(12.2), Inches(0.3),
             "What makes it different / innovative?", 12, True, SAFFRON)
    add_text(s, Inches(0.55), Inches(6.25), Inches(12.2), Inches(0.7),
             "82 features across 11 modules in one shield • One AI Assistant that operates prediction, maps, health, agri, and government workflows • Unique extras: WhatsApp alert bot, siren mode, family reunion board, fake-news filter, priority support for women/children, clean-water & power-outage maps • Software-only, Bharat-first, offline-ready.",
             13, False, WHITE)

    # ========== SLIDE 6: TECH STACK ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "TECH STACK & ARCHITECTURE")

    stacks = [
        ("Frontend", ["Next.js", "React", "TypeScript", "Tailwind CSS", "PWA"]),
        ("Backend / API", ["Next.js Route Handlers", "FastAPI", "Python", "Prisma"]),
        ("Data / Security", ["Supabase", "PostgreSQL", "Supabase Auth", "RBAC"]),
        ("AI / Maps / Deploy", ["Generative AI", "ML / CV / NLP", "Mapbox", "Vercel + GitHub"]),
    ]
    for i, (title, items) in enumerate(stacks):
        left = Inches(0.35 + i * 3.2)
        add_card(s, left, Inches(1.3), Inches(3.05), Inches(3.0), WHITE)
        add_text(s, left + Inches(0.15), Inches(1.45), Inches(2.75), Inches(0.4), title, 14, True, SAFFRON)
        add_bullets(s, left + Inches(0.15), Inches(1.95), Inches(2.75), Inches(2.2), items, 13)

    add_card(s, Inches(0.35), Inches(4.5), Inches(6.2), Inches(2.6), WHITE)
    add_text(s, Inches(0.55), Inches(4.65), Inches(5.8), Inches(0.35), "APIs / Libraries", 13, True, SAFFRON)
    add_bullets(s, Inches(0.55), Inches(5.1), Inches(5.8), Inches(1.8), [
        "Weather / rainfall APIs for multi-hazard risk signals",
        "Mapbox / GIS for live maps, routes, and geolocation",
        "LLM APIs for multilingual assistant + scheme guidance",
        "Vision models for crop/building/flood damage from images",
        "STT / TTS for voice interaction in Indian languages",
    ], 12)

    add_card(s, Inches(6.8), Inches(4.5), Inches(6.15), Inches(2.6), WHITE)
    add_text(s, Inches(7.0), Inches(4.65), Inches(5.8), Inches(0.35), "Deployment Plan", 13, True, SAFFRON)
    add_bullets(s, Inches(7.0), Inches(5.1), Inches(5.8), Inches(1.8), [
        "Frontend + APIs on Vercel",
        "Auth + Postgres on Supabase",
        "AI/ML services via FastAPI (Python)",
        "Repo + CI on GitHub",
        "PWA install for offline emergency basics",
    ], 12)

    # ========== SLIDE 7: ARCHITECTURE DIAGRAM ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "ARCHITECTURE DIAGRAM")

    add_text(s, Inches(0.4), Inches(1.2), Inches(12.5), Inches(0.35),
             "High-level Technical Flow — BharatRaksha AI", 14, True, NAVY, PP_ALIGN.CENTER)

    # Row 1 users
    users = [("Citizen App", TEAL), ("AI Assistant", SAFFRON), ("Govt / Rescue Dashboard", BLUE)]
    for i, (label, color) in enumerate(users):
        flow_box(s, Inches(1.2 + i * 3.8), Inches(1.7), Inches(3.2), Inches(0.7), label, color)

    # arrows down
    for i in range(3):
        arrow_down(s, Inches(2.6 + i * 3.8), Inches(2.5))

    flow_box(s, Inches(3.5), Inches(2.9), Inches(6.3), Inches(0.65),
             "Next.js Frontend  +  PWA Offline Layer", NAVY)

    arrow_down(s, Inches(6.45), Inches(3.65))

    # middle APIs
    mids = [("Next.js APIs", BLUE), ("FastAPI AI Layer", TEAL), ("Mapbox / Weather APIs", SAFFRON)]
    for i, (label, color) in enumerate(mids):
        flow_box(s, Inches(1.0 + i * 4.0), Inches(4.05), Inches(3.5), Inches(0.65), label, color)

    for i in range(3):
        arrow_down(s, Inches(2.55 + i * 4.0), Inches(4.8))

    # bottom data/ai
    bottoms = [("Supabase Auth + Postgres", NAVY), ("ML / CV / NLP Models", RED), ("Secure Cloud Backup", GREEN)]
    for i, (label, color) in enumerate(bottoms):
        flow_box(s, Inches(1.0 + i * 4.0), Inches(5.2), Inches(3.5), Inches(0.65), label, color)

    add_text(s, Inches(0.4), Inches(6.15), Inches(12.5), Inches(0.9),
             "Flow: User / Authority → Frontend → APIs → AI Models + Maps/Weather → Database → Live alerts, guidance, dashboards & recovery support",
             13, False, GRAY, PP_ALIGN.CENTER)

    # ========== SLIDE 8: FEATURE FLOWCHART ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, LIGHT)
    add_header_bar(s, "END-TO-END FEATURE FLOW")

    add_text(s, Inches(0.4), Inches(1.2), Inches(12.5), Inches(0.35),
             "How BharatRaksha AI works in a real disaster", 14, True, NAVY, PP_ALIGN.CENTER)

    steps = [
        ("1. Sense", "Weather + reports\n+ risk signals", BLUE),
        ("2. Predict", "AI multi-hazard\nrisk score", TEAL),
        ("3. Alert", "Multilingual AI\nwarnings & guidance", SAFFRON),
        ("4. Respond", "Maps, routes,\nshelters, hospitals", RED),
        ("5. Coordinate", "Govt dashboard\nresources & teams", NAVY),
        ("6. Recover", "Health + agri\ndamage & schemes", GREEN),
    ]
    for i, (t, d, c) in enumerate(steps):
        left = Inches(0.35 + i * 2.15)
        flow_box(s, left, Inches(1.8), Inches(2.0), Inches(0.55), t, c)
        add_card(s, left, Inches(2.55), Inches(2.0), Inches(1.3), WHITE)
        add_text(s, left + Inches(0.08), Inches(2.7), Inches(1.85), Inches(1.0), d, 11, False, DARK, PP_ALIGN.CENTER)
        if i < 5:
            arrow_right(s, left + Inches(2.0), Inches(2.0))

    # modules overview
    add_text(s, Inches(0.4), Inches(4.1), Inches(12.5), Inches(0.35),
             "82 Features across 11 Modules", 14, True, NAVY)

    modules = [
        "1 Prediction", "2 Maps", "3 Emergency", "4 AI Assistant",
        "5 Healthcare", "6 Agriculture", "7 Computer Vision",
        "8 Govt Dashboard", "9 Community", "10 Platform", "11 Extra Tools",
    ]
    for i, m in enumerate(modules):
        left = Inches(0.35 + (i % 6) * 2.15)
        top = Inches(4.55) if i < 6 else Inches(5.45)
        if i == 10:
            left = Inches(0.35)
            top = Inches(6.35)
        flow_box(s, left, top, Inches(2.0), Inches(0.55), m, NAVY if i < 10 else SAFFRON)

    if True:
        add_text(s, Inches(2.5), Inches(6.4), Inches(10), Inches(0.5),
                 "Extra: WhatsApp bot • Siren mode • Family reunion • Fake-news filter • Clean water / power maps • Priority support",
                 12, False, GRAY)

    # ========== SLIDE 9: THANK YOU ==========
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(2.7), Inches(2.9), Inches(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = SAFFRON
    line.line.fill.background()

    add_text(s, Inches(1), Inches(2.0), Inches(11.3), Inches(0.7),
             "THANK YOU", 48, True, WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(3.0), Inches(11.3), Inches(0.5),
             "BharatRaksha AI — The AI Shield for Every Disaster", 20, False, SAFFRON, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(3.8), Inches(11.3), Inches(0.8),
             "Predict early. Respond faster. Recover stronger.", 16, False, WHITE, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(5.0), Inches(11.3), Inches(0.4),
             "Team: [ Team Name ]   |   Decode SIH 2026   |   OSCode by communityX", 13, False, RGBColor(0xB8, 0xC4, 0xD6), PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(5.6), Inches(11.3), Inches(0.4),
             "Contact: [ email / phone ]", 13, False, RGBColor(0x9A, 0xA8, 0xBC), PP_ALIGN.CENTER)

    out = "/Users/rohith_kumar_505/BharatRaksha-AI/BharatRaksha_AI_Decode_SIH_2026_Ideation.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
